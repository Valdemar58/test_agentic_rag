"""Генераторы минимальных файлов для тестов и режима самопроверки: PDF с текстовым слоем
и без него, DOCX (с таблицей и разделом терминов по желанию), изображение-«скан»."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable

from docx import Document
from PIL import Image

ALTCHUNK_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/aFChunk"
WML_MAIN_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
CONTENT_TYPES_PART = "[Content_Types].xml"
DOCUMENT_RELS_PART = "word/_rels/document.xml.rels"
DOCUMENT_PART = "word/document.xml"
# Время записей zip фиксировано: zip хранит mtime с шагом 2 с, и одинаковые docx, собранные в разные
# секунды, иначе различались бы байтами и sha256 (тесты инкрементальности плавали из-за этого)
FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
FILE_MODE_REGULAR = 0o644 << 16


def _fixed_info(name: str) -> zipfile.ZipInfo:
    info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = FILE_MODE_REGULAR
    return info


def deterministic_zip(data: bytes, part_edits: dict[str, Callable[[bytes], bytes]] | None = None) -> bytes:
    """Перепаковывает zip (docx, xlsx) с фиксированным временем записей: байты воспроизводимы.

    `part_edits` — правки отдельных частей по имени (например, дата изменения в docProps/core.xml)."""
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            blob = source.read(info.filename)
            edit = (part_edits or {}).get(info.filename)
            target.writestr(_fixed_info(info.filename), edit(blob) if edit else blob)
    return buffer.getvalue()


def minimal_pdf_bytes(text: str | None) -> bytes:
    """Одностраничный PDF. При text=None страница пустая — имитация скана без текстового слоя."""
    if text is None:
        content = b""
    else:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1", errors="replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    buffer = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(buffer))
        buffer += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(buffer)
    buffer += f"xref\n0 {len(objects) + 1}\n".encode()
    buffer += b"0000000000 65535 f \n"
    for offset in offsets:
        buffer += f"{offset:010d} 00000 n \n".encode()
    buffer += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(buffer)


def minimal_docx_bytes(
    paragraphs: list[str], *, with_table: bool = False, terms_heading: str | None = None
) -> bytes:
    """DOCX с абзацами; опционально таблица 2×2 и заголовок раздела терминов."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if terms_heading:
        document.add_heading(terms_heading, level=1)
        document.add_paragraph("СИЗ — средства индивидуальной защиты.")
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Показатель"
        table.cell(0, 1).text = "Значение"
        table.cell(1, 0).text = "Норма"
        table.cell(1, 1).text = "1"
    buffer = io.BytesIO()
    document.save(buffer)
    return deterministic_zip(buffer.getvalue())


def _rewrite_docx(data: bytes, edits: dict[str, bytes], extra: dict[str, bytes]) -> bytes:
    """Пересобирает zip-пакет DOCX: части из edits заменяются, части из extra добавляются."""
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            target.writestr(_fixed_info(info.filename), edits.get(info.filename, source.read(info.filename)))
        for name, blob in extra.items():
            target.writestr(_fixed_info(name), blob)
    return buffer.getvalue()


def docx_with_misdeclared_altchunk_bytes(paragraphs: list[str], *, with_table: bool = False) -> bytes:
    """DOCX, как его делает шаблонизатор Тессы для файлов «Для печати_…»: вложенный altChunk-docx
    объявлен в [Content_Types].xml как XML-часть главного документа. Word открывает, python-docx — нет
    (реальный случай из экспорта заказчика 2026-09-14)."""
    base = minimal_docx_bytes(paragraphs, with_table=with_table)
    with zipfile.ZipFile(io.BytesIO(base)) as package:
        content_types = package.read(CONTENT_TYPES_PART).decode("utf-8")
        rels = package.read(DOCUMENT_RELS_PART).decode("utf-8")
    content_types = content_types.replace(
        "</Types>", f'<Default Extension="docx" ContentType="{WML_MAIN_CONTENT_TYPE}"/></Types>'
    )
    rels = rels.replace(
        "</Relationships>",
        f'<Relationship Id="AltChunkId1" Type="{ALTCHUNK_REL_TYPE}" Target="/word/afchunk1.docx"/>'
        "</Relationships>",
    )
    return _rewrite_docx(
        base,
        {CONTENT_TYPES_PART: content_types.encode("utf-8"), DOCUMENT_RELS_PART: rels.encode("utf-8")},
        {"word/afchunk1.docx": minimal_docx_bytes(["Блок электронной подписи"])},
    )


def docx_with_broken_part_bytes() -> bytes:
    """DOCX с испорченным word/document.xml — действительно битый файл."""
    return _rewrite_docx(minimal_docx_bytes(["x"]), {DOCUMENT_PART: b"\x00\x01 not xml at all"}, {})


def minimal_image_bytes(image_format: str = "JPEG") -> bytes:
    """Маленькое изображение — имитация скана."""
    image = Image.new("RGB", (32, 32), color=(255, 255, 255))
    buffer = io.BytesIO()
    image.save(buffer, format=image_format)
    return buffer.getvalue()
