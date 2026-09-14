"""Проверка скачанных файлов лёгкими библиотеками (§8.3): открывается ли файл, есть ли
текстовый слой, таблицы и раздел терминов. Для сканов признаки таблиц и терминов — «н/д»."""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from uuid import UUID

from tessa_export.config import CoverageSettings, FilesSettings
from tessa_export.files import FileRecord

logger = logging.getLogger(__name__)

MIN_TEXT_CHARS_PER_PAGE = 20
MAX_PDF_PAGES_TO_SCAN = 30
MAX_HEADING_LENGTH = 120


@dataclass(frozen=True)
class FileInspection:
    smoke_ok: bool
    smoke_error: str | None = None
    has_text_layer: bool | None = None
    page_count: int | None = None
    has_tables: bool | None = None
    has_terms_section: bool | None = None


def _terms_found(lines: list[str], patterns: list[re.Pattern[str]]) -> bool:
    for line in lines:
        candidate = line.strip()
        if not candidate or len(candidate) > MAX_HEADING_LENGTH:
            continue
        if any(pattern.search(candidate) for pattern in patterns):
            return True
    return False


def _inspect_pdf(data: bytes, patterns: list[re.Pattern[str]]) -> FileInspection:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    page_count = len(reader.pages)
    texts: list[str] = []
    for page in reader.pages[:MAX_PDF_PAGES_TO_SCAN]:
        texts.append(page.extract_text() or "")
    text_pages = sum(1 for text in texts if len(text.strip()) >= MIN_TEXT_CHARS_PER_PAGE)
    has_text_layer = text_pages > 0
    lines = [line for text in texts for line in text.splitlines()]
    return FileInspection(
        smoke_ok=True,
        has_text_layer=has_text_layer,
        page_count=page_count,
        has_tables=None,
        has_terms_section=_terms_found(lines, patterns) if has_text_layer else None,
    )


def _inspect_docx(data: bytes, patterns: list[re.Pattern[str]]) -> FileInspection:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    return FileInspection(
        smoke_ok=True,
        has_text_layer=True,
        has_tables=len(document.tables) > 0,
        has_terms_section=_terms_found(lines, patterns),
    )


def _inspect_xlsx(data: bytes) -> FileInspection:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True)
    sheet_count = len(workbook.sheetnames)
    workbook.close()
    return FileInspection(smoke_ok=True, has_text_layer=True, page_count=sheet_count, has_tables=True)


def _inspect_pptx(data: bytes, patterns: list[re.Pattern[str]]) -> FileInspection:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    lines: list[str] = []
    has_tables = False
    slide_count = 0
    for slide in presentation.slides:
        slide_count += 1
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                has_tables = True
            if getattr(shape, "has_text_frame", False):
                lines.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
    return FileInspection(
        smoke_ok=True,
        has_text_layer=True,
        page_count=slide_count,
        has_tables=has_tables,
        has_terms_section=_terms_found(lines, patterns),
    )


def _inspect_image(data: bytes) -> FileInspection:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as image:
        image.verify()
    return FileInspection(smoke_ok=True, has_text_layer=False, page_count=1)


def inspect_file(
    data: bytes, extension: str, image_extensions: set[str], terms_patterns: list[str]
) -> FileInspection:
    patterns = [re.compile(pattern, re.IGNORECASE) for pattern in terms_patterns]
    try:
        if extension == "pdf":
            return _inspect_pdf(data, patterns)
        if extension == "docx":
            return _inspect_docx(data, patterns)
        if extension == "xlsx":
            return _inspect_xlsx(data)
        if extension == "pptx":
            return _inspect_pptx(data, patterns)
        if extension in image_extensions:
            return _inspect_image(data)
        return FileInspection(smoke_ok=True)
    except Exception as exc:  # noqa: BLE001 — любая ошибка парсинга должна попасть в отчёт, а не уронить экспорт
        return FileInspection(smoke_ok=False, smoke_error=f"{type(exc).__name__}: {exc}")


def inspect_records(
    export_root: Path,
    file_records: dict[UUID, list[FileRecord]],
    files_settings: FilesSettings,
    coverage: CoverageSettings,
) -> None:
    """Заполняет признаки у скачанных файлов; пропущенные файлы не трогает."""
    image_extensions = set(files_settings.image_extensions)
    for card_id, records in file_records.items():
        for record in records:
            if not record.downloaded or record.relative_path is None:
                continue
            data = (export_root / record.relative_path).read_bytes()
            inspection = inspect_file(
                data, record.file.extension, image_extensions, coverage.terms_section_patterns
            )
            record.smoke_ok = inspection.smoke_ok
            record.smoke_error = inspection.smoke_error
            record.has_text_layer = inspection.has_text_layer
            record.page_count = inspection.page_count
            record.has_tables = inspection.has_tables
            record.has_terms_section = inspection.has_terms_section
            if not inspection.smoke_ok:
                logger.warning(
                    "Файл «%s» карточки %s не открывается: %s",
                    record.file.name,
                    card_id,
                    inspection.smoke_error,
                )
