"""Синтетические документы с русским текстом.

DOCX со структурой заголовков, нумерованными пунктами и таблицами (нативный путь Docling), PDF с
текстовым слоем (fpdf2 + шрифт DejaVu Sans, кириллица), «сканы» — изображения страниц и PDF без
текстового слоя (путь VLM). Формулировки условные и не относятся к реальной организации.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import UTC, datetime

from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from common.config import ROOT
from tessa_export.sample_files import deterministic_zip

FONT_PATH = ROOT / "assets" / "fonts" / "DejaVuSans.ttf"
FONT_FAMILY = "DejaVu"
SCAN_PAGE_SIZE = (1240, 1754)  # A4 при 150 dpi
SCAN_MARGIN = 100
SCAN_FONT_SIZE = 28
SCAN_LINE_HEIGHT = 42
SCAN_TITLE_FONT_SIZE = 36
PDF_TITLE_SIZE = 14
PDF_HEADING_SIZE = 12
PDF_BODY_SIZE = 10
PDF_LINE_HEIGHT = 6
TABLE_CELL_SEPARATOR = " | "
# Даты в свойствах xlsx фиксированы: иначе openpyxl пишет «сейчас» (modified — прямо при сохранении),
# и одинаковые файлы различаются байтами и sha256
FIXED_DOCUMENT_TIME = datetime(2026, 1, 1, tzinfo=UTC)
CORE_PROPERTIES_PART = "docProps/core.xml"
_MODIFIED_RE = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _fix_modified(core_xml: bytes) -> bytes:
    stamp = FIXED_DOCUMENT_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    return _MODIFIED_RE.sub(lambda match: match.group(1) + stamp + match.group(2), core_xml)


@dataclass(frozen=True)
class Section:
    """Раздел документа: заголовок, абзацы (пункты) и необязательная таблица (первая строка — шапка)."""

    title: str
    paragraphs: tuple[str, ...] = ()
    table: tuple[tuple[str, ...], ...] | None = None
    level: int = 1


@dataclass(frozen=True)
class DocumentText:
    title: str
    preamble: tuple[str, ...] = ()
    sections: tuple[Section, ...] = ()

    def lines(self) -> list[str]:
        """Плоский текст документа построчно — для сканов и проверок."""
        lines = [self.title, *self.preamble]
        for section in self.sections:
            lines.append(section.title)
            lines.extend(section.paragraphs)
            if section.table:
                lines.extend(TABLE_CELL_SEPARATOR.join(row) for row in section.table)
        return lines


def docx_bytes(text: DocumentText) -> bytes:
    """DOCX: заголовок документа — Heading 1, разделы — Heading 2/3, таблицы со стилем Table Grid."""
    document = Document()
    document.add_heading(text.title, level=1)
    for paragraph in text.preamble:
        document.add_paragraph(paragraph)
    for section in text.sections:
        document.add_heading(section.title, level=min(section.level + 1, 4))
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
        if section.table:
            rows, cols = len(section.table), len(section.table[0])
            table = document.add_table(rows=rows, cols=cols)
            table.style = "Table Grid"
            for row_index, row in enumerate(section.table):
                for col_index, value in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.text = value
                    if row_index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
    buffer = io.BytesIO()
    document.save(buffer)
    return deterministic_zip(buffer.getvalue())


def pdf_bytes(text: DocumentText) -> bytes:
    """PDF с текстовым слоем (кириллица через встроенный DejaVu Sans)."""
    pdf = FPDF()
    pdf.add_font(FONT_FAMILY, style="", fname=str(FONT_PATH))
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    def line(content: str, *, size: int, height: float) -> None:
        pdf.set_font(FONT_FAMILY, size=size)
        pdf.multi_cell(0, height, content, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    line(text.title, size=PDF_TITLE_SIZE, height=PDF_LINE_HEIGHT + 2)
    pdf.ln(2)
    for paragraph in text.preamble:
        line(paragraph, size=PDF_BODY_SIZE, height=PDF_LINE_HEIGHT)
    for section in text.sections:
        pdf.ln(2)
        line(section.title, size=PDF_HEADING_SIZE, height=PDF_LINE_HEIGHT + 1)
        for paragraph in section.paragraphs:
            line(paragraph, size=PDF_BODY_SIZE, height=PDF_LINE_HEIGHT)
        if section.table:
            for row in section.table:
                line(TABLE_CELL_SEPARATOR.join(row), size=PDF_BODY_SIZE, height=PDF_LINE_HEIGHT)
    return bytes(pdf.output())


def _wrap(draw: ImageDraw.ImageDraw, line: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    words = line.split()
    wrapped: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if current and draw.textlength(candidate, font=font) > width:
            wrapped.append(current)
            current = word
        else:
            current = candidate
    if current:
        wrapped.append(current)
    return wrapped or [""]


def _render_page(lines: list[str], *, title: str | None) -> Image.Image:
    image = Image.new("RGB", SCAN_PAGE_SIZE, color=(250, 248, 243))
    draw = ImageDraw.Draw(image)
    body_font = ImageFont.truetype(str(FONT_PATH), SCAN_FONT_SIZE)
    title_font = ImageFont.truetype(str(FONT_PATH), SCAN_TITLE_FONT_SIZE)
    width = SCAN_PAGE_SIZE[0] - 2 * SCAN_MARGIN
    y = SCAN_MARGIN
    if title:
        for piece in _wrap(draw, title, title_font, width):
            draw.text((SCAN_MARGIN, y), piece, fill=(20, 20, 20), font=title_font)
            y += SCAN_LINE_HEIGHT + 8
        y += SCAN_LINE_HEIGHT // 2
    for line in lines:
        for piece in _wrap(draw, line, body_font, width):
            if y > SCAN_PAGE_SIZE[1] - SCAN_MARGIN:
                break
            draw.text((SCAN_MARGIN, y), piece, fill=(30, 30, 30), font=body_font)
            y += SCAN_LINE_HEIGHT
        y += SCAN_LINE_HEIGHT // 3
    return image


def scan_image_bytes(text: DocumentText, *, image_format: str = "JPEG") -> bytes:
    """Одна страница «скана» с русским текстом (JPEG/PNG), без текстового слоя."""
    image = _render_page(text.lines()[1:], title=text.title)
    buffer = io.BytesIO()
    image.save(buffer, format=image_format, quality=85)
    return buffer.getvalue()


def scan_pdf_bytes(text: DocumentText, *, lines_per_page: int = 28) -> bytes:
    """PDF из изображений страниц — скан без текстового слоя (маршрут dots.mocr)."""
    lines = text.lines()[1:]
    pages = [lines[i : i + lines_per_page] for i in range(0, max(len(lines), 1), lines_per_page)]
    images = [
        _render_page(page, title=text.title if index == 0 else None).convert("RGB")
        for index, page in enumerate(pages)
    ]
    buffer = io.BytesIO()
    images[0].save(buffer, format="PDF", save_all=True, append_images=images[1:], resolution=150.0)
    return buffer.getvalue()


def xlsx_bytes(sheet_title: str, rows: tuple[tuple[str, ...], ...]) -> bytes:
    """Таблица XLSX: первая строка — шапка."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = sheet_title[:31]
    for row in rows:
        sheet.append(list(row))
    workbook.properties.created = FIXED_DOCUMENT_TIME
    buffer = io.BytesIO()
    workbook.save(buffer)
    return deterministic_zip(buffer.getvalue(), part_edits={CORE_PROPERTIES_PART: _fix_modified})
