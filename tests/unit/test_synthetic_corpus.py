"""Синтетический корпус: документы читаются, экспорт валиден (2.6), реальный корпус не затирается."""

from __future__ import annotations

import io
import json
import logging
from collections.abc import Iterator
from pathlib import Path

import pytest
import structlog
from docx import Document
from PIL import Image
from pypdf import PdfReader

from synthetic import texts
from synthetic.corpus import (
    ARCHIVE_NAME,
    README_NAME,
    RealCorpusPresentError,
    build_gateway,
    generate_corpus,
)
from synthetic.documents import docx_bytes, pdf_bytes, scan_image_bytes, scan_pdf_bytes, xlsx_bytes


@pytest.fixture(autouse=True)
def _reset_logging() -> Iterator[None]:
    yield
    root = logging.getLogger()
    for handler in list(root.handlers):
        root.removeHandler(handler)
        handler.close()
    structlog.reset_defaults()


def test_docx_has_headings_paragraphs_and_table() -> None:
    document = Document(io.BytesIO(docx_bytes(texts.REGULATION_OT)))
    headings = [
        p.text for p in document.paragraphs if p.style is not None and p.style.name.startswith("Heading")
    ]
    assert headings[0] == "Положение об охране труда"
    assert "2. Термины и определения" in headings
    assert any("СИЗ — средства индивидуальной защиты" in p.text for p in document.paragraphs)
    assert len(document.tables) == 1 and document.tables[0].cell(0, 0).text == "Вид инструктажа"


def test_pdf_has_cyrillic_text_layer_and_scan_has_none() -> None:
    text = "\n".join(
        page.extract_text() for page in PdfReader(io.BytesIO(pdf_bytes(texts.CONTRACT_D1))).pages
    )
    assert "Договор поставки" in text and "1 200 000" in text
    scan = PdfReader(io.BytesIO(scan_pdf_bytes(texts.LETTER_IN_33)))
    assert len(scan.pages) >= 1
    assert all(not page.extract_text().strip() for page in scan.pages)
    image = Image.open(io.BytesIO(scan_image_bytes(texts.INSTRUCTION_FIRE)))
    assert image.size == (1240, 1754) and image.format == "JPEG"
    assert len(xlsx_bytes("Лист", texts.CONTRACT_SPECIFICATION)) > 0


def test_gateway_scenario_has_all_cards_and_files() -> None:
    gateway = build_gateway()
    assert len(gateway.cards) == 13
    for snapshot in gateway.cards.values():
        for file in snapshot.files:
            if not file.is_virtual:
                assert (snapshot.card_id, file.row_id) in gateway.contents, file.name


def test_generate_corpus_passes_export_validation(tmp_path: Path) -> None:
    summary = generate_corpus(tmp_path / "corpus")
    assert summary.overall == "PASS"
    assert summary.documents == 13
    manifest = json.loads(summary.manifest_path.read_text(encoding="utf-8"))
    assert manifest["synthetic"] is True
    statuses = manifest["stats"]["doc_statuses"]
    assert statuses.get("cancelled") == 1 and statuses.get("draft") == 1 and statuses.get("active") == 11
    assert manifest["stats"]["doc_kinds"]["Приказ"] == 6
    relation_types = manifest["stats"]["relation_types"]
    for relation in ("в отмену", "дополнение", "Приказ", "Основной договор", "запрос"):
        assert relation in relation_types, relation
    graph = json.loads(summary.links_graph_path.read_text(encoding="utf-8"))
    assert any(edge["relation_type"] == "в отмену" for edge in graph["edges"])

    files = [file for doc in manifest["documents"] for file in doc["files"] if file["downloaded"]]
    assert any(file["has_text_layer"] is False and file["extension"] == "pdf" for file in files)
    assert any(file["has_text_layer"] is True and file["extension"] == "pdf" for file in files)
    assert any(file["has_terms_section"] is True and file["extension"] == "docx" for file in files)
    assert any(file["smoke_note"] for file in files if file["name"].startswith("Для печати_"))
    assert any(
        file["skipped_reason"]
        for doc in manifest["documents"]
        for file in doc["files"]
        if file["name"].endswith(".sig")
    )

    report = summary.report_path.read_text(encoding="utf-8")
    assert "синтетические" in report.lower()
    assert (
        (tmp_path / "corpus" / README_NAME).read_text(encoding="utf-8").startswith("# ДАННЫЕ СИНТЕТИЧЕСКИЕ")
    )
    assert (tmp_path / "corpus" / ARCHIVE_NAME).is_file()


def test_real_export_is_not_overwritten_without_force(tmp_path: Path) -> None:
    export = tmp_path / "corpus" / "export"
    export.mkdir(parents=True)
    (export / "manifest.json").write_text('{"synthetic": false, "documents": []}', encoding="utf-8")
    with pytest.raises(RealCorpusPresentError):
        generate_corpus(tmp_path / "corpus")
    summary = generate_corpus(tmp_path / "corpus", force=True)
    assert summary.overall == "PASS"
