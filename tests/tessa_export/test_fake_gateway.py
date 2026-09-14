"""Фейковый шлюз, конструктор снимков и генераторы файлов (основа AC-EXP.1)."""

from __future__ import annotations

import io
from typing import Any
from uuid import UUID, uuid4

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader

from tessa_export.fake import FakeGateway, build_demo_scenario, link, make_file, make_snapshot, stable_uuid
from tessa_export.models import (
    COMMON_SECTION,
    INCOMING_SECTION,
    OUTGOING_SECTION,
    CardNotFoundError,
    GatewayConnectionError,
    parse_links,
)
from tessa_export.sample_files import minimal_docx_bytes, minimal_image_bytes, minimal_pdf_bytes


def test_snapshot_structure_matches_real_card_shape() -> None:
    target = stable_uuid("card", "X")
    snapshot = make_snapshot(uuid4(), outgoing=[link(target)], files=[make_file(uuid4(), "Приказ.docx")])
    assert snapshot.common_text("DocTypeTitle") == "Приказ"
    assert snapshot.common_text("DepartmentName")
    assert snapshot.outgoing[0].doc_id == target
    assert snapshot.outgoing[0].ref_type_name == "в отмену"
    raw_card = snapshot.raw["Card"]
    assert raw_card["ID::uid"] == str(snapshot.card_id)
    assert raw_card["Sections"][OUTGOING_SECTION]["Rows"][0]["DocID::uid"] == str(target)
    assert raw_card["Sections"][COMMON_SECTION]["Fields"]["StatusID::uid"]
    assert snapshot.card_data_json["sections"][COMMON_SECTION]["fields"]["FullNumber"] == "1"
    assert snapshot.card_data_json["files"][0]["name"] == "Приказ.docx"
    assert snapshot.files[0].extension == "docx"


def test_parse_links_accepts_str_and_uuid_and_skips_empty() -> None:
    doc_id = uuid4()
    rows: list[dict[str, Any]] = [
        {"DocID": str(doc_id), "RefTypeName": "в отмену", "Order": 1},
        {"DocID": doc_id, "DocDescription": "x"},
        {"DocID": None},
        {"DocID": "not-a-uuid"},
    ]
    links = parse_links(rows)
    assert [item.doc_id for item in links] == [doc_id, doc_id]
    assert links[0].ref_type_name == "в отмену"
    assert links[0].order == 1
    assert parse_links(None) == []


def test_fake_gateway_behaviour() -> None:
    gateway = FakeGateway()
    card_id = uuid4()
    file = make_file(card_id, "Документ.pdf")
    gateway.add(make_snapshot(card_id, files=[file]), {"Документ.pdf": b"%PDF"})
    snapshot = gateway.get_card(card_id)
    assert snapshot.card_id == card_id
    downloaded = gateway.download_file(card_id, file)
    assert downloaded.content == b"%PDF"
    assert downloaded.file_name == "Документ.pdf"
    assert downloaded.content_type == "application/pdf"
    unknown = uuid4()
    with pytest.raises(CardNotFoundError):
        gateway.get_card(unknown)
    broken = uuid4()
    gateway.card_errors[broken] = GatewayConnectionError("сеть")
    with pytest.raises(GatewayConnectionError):
        gateway.get_card(broken)
    assert gateway.get_calls == [card_id, unknown, broken]
    gateway.close()
    assert gateway.closed


def test_demo_scenario_covers_required_cases() -> None:
    gateway, seed = build_demo_scenario()
    assert len(seed) == 2
    a = gateway.get_card(seed[0])
    b_id = a.outgoing[0].doc_id
    b = gateway.get_card(b_id)
    # цикл A → B → A
    assert any(item.doc_id == a.card_id for item in b.outgoing)
    # входящие без типа
    assert a.incoming[0].ref_type_name is None
    assert a.sections[INCOMING_SECTION].rows is not None
    # виртуальный, неподдерживаемый и кириллический файлы
    assert any(item.is_virtual for item in a.files)
    assert any(item.extension == "sig" for item in a.files)
    assert any(item.name.startswith("Для печати") for item in a.files)
    # карточка без файлов и с двумя путями входа
    e_id = a.outgoing[1].doc_id
    e = gateway.get_card(e_id)
    assert e.files == []
    assert len(e.incoming) == 2
    # цепочка до глубины 3 от A: A → B → C → D
    c = gateway.get_card(next(item.doc_id for item in b.outgoing if item.doc_id != a.card_id))
    d = gateway.get_card(c.outgoing[0].doc_id)
    assert d.files[0].extension == "pdf"
    # содержимое файлов доступно для всех невиртуальных файлов
    for card in gateway.cards.values():
        for file in card.files:
            if not file.is_virtual:
                assert gateway.download_file(card.card_id, file).content


def test_sample_files_are_parseable() -> None:
    text_pdf = PdfReader(io.BytesIO(minimal_pdf_bytes("Hello order")))
    assert "Hello order" in (text_pdf.pages[0].extract_text() or "")
    scan_pdf = PdfReader(io.BytesIO(minimal_pdf_bytes(None)))
    assert not (scan_pdf.pages[0].extract_text() or "").strip()
    document = Document(io.BytesIO(minimal_docx_bytes(["Абзац"], with_table=True, terms_heading="Термины")))
    assert len(document.tables) == 1
    assert any("Термины" in paragraph.text for paragraph in document.paragraphs)
    image = Image.open(io.BytesIO(minimal_image_bytes()))
    assert image.size == (32, 32)


def test_stable_uuid_is_deterministic() -> None:
    assert stable_uuid("card", "A") == stable_uuid("card", "A")
    assert stable_uuid("card", "A") != stable_uuid("card", "B")
    assert isinstance(stable_uuid("x"), UUID)
